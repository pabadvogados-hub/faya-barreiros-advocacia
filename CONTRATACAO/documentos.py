"""
Contrato, procuracao e declaracao preenchidos a partir da qualificacao e da triagem.
Os modelos ficam em DOCS_MODELOS/ com {{CAMPOS}} (lista em docs/CAMPOS_DOS_MODELOS.md).
"""
import json
import os
import re
import sys

sys.path.insert(0, os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), 'NUCLEO'))
import ambiente  # noqa: E402,F401  (carrega .env e caminhos)
from datetime import date, datetime, timedelta

from docx import Document

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from docx_escritorio import preencher_modelo  # noqa: E402
from pasta_cliente import CONTRATACAO, nome_pasta, raiz_clientes  # noqa: E402
from relatorio_triagem import endereco_completo  # noqa: E402

RAIZ = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, RAIZ)
from config.escritorio import ESCRITORIO, HONORARIOS, OUTORGADOS, PRAZOS, SOCIO, TITULAR  # noqa: E402

MODELOS = os.path.join(RAIZ, 'DOCS_MODELOS')
# Modelos oficiais do escritorio (enviados em 18/10/2026), por area. Contrato, procuracao e declaracao
# da Dra. Faya (previdenciario) e contrato de gestao do passivo do Dr. William (tributario).
DOCUMENTOS_POR_AREA = {
    'PREVIDENCIARIO': [
        ('Contrato de Honorarios', 'PREV_CONTRATO_EXITO.docx'),
        ('Procuracao', 'PROCURACAO_PREVIDENCIARIA_MODELO.docx'),
        ('Declaracao de Hipossuficiencia', 'DECLARACAO_HIPOSSUFICIENCIA_MODELO.docx'),
    ],
    'TRIBUTARIO': [
        ('Contrato de Honorarios', 'TRIB_CONTRATO_GESTAO_PASSIVO.docx'),
        ('Procuracao', 'PROCURACAO_MODELO.docx'),   # PROVISORIA: o Dr. William ainda nao mandou a dele
    ],
}
# Variantes escolhidas pelo cadastro (chave: valor no CADASTRO.txt).
CONTRATO_EXITO_SALARIOS = 'PREV_CONTRATO_EXITO_SALARIOS.docx'   # "Contrato: exito e salarios" (cobra R$ 300 iniciais)
CONTRATO_TRIB_COM_RESULTADO = 'TRIB_CONTRATO_GESTAO_PASSIVO_COM_RESULTADO.docx'   # "Contrato: com resultado"
TERMO_RENUNCIA = ('Termo de Renuncia (JEF)', 'PREV_TERMO_RENUNCIA_JEF.docx')       # "Juizado federal: sim"
MESES = ['janeiro', 'fevereiro', 'marco', 'abril', 'maio', 'junho', 'julho', 'agosto', 'setembro',
         'outubro', 'novembro', 'dezembro']
PENDENCIA = re.compile(r'\[(?:PREENCHER|CONFERIR)[^\]]*\]')


def data_extenso(d=None):
    d = d or date.today()
    mes = MESES[d.month - 1].replace('marco', 'março')
    return f'{d.day} de {mes} de {d.year}'


def _data(txt):
    try:
        return datetime.strptime((txt or '').strip()[:10], '%d/%m/%Y').date()
    except ValueError:
        return None


def data_contrato(cadastro, triagem):
    return (_data(cadastro.get('data_contrato') or cadastro.get('data_do_contrato'))
            or _data(triagem.get('data_reuniao')) or date.today())


def prazos_do_caso(inicio):
    """Marcos da contratacao ate o protocolo da inicial (config/escritorio.py PRAZOS)."""
    p = PRAZOS
    documentos = inicio + timedelta(days=p['documentos_dias_apos_contrato'])
    minuta = documentos + timedelta(days=p['inicial_dias_apos_documentos'])
    fmt = lambda d: d.strftime('%d/%m/%Y')  # noqa: E731
    return [
        {'marco': 'Cobrar do cliente os documentos do checklist', 'id': 'documentos',
         'data': fmt(documentos), 'responsavel': 'Comercial'},
        {'marco': 'Minuta da inicial pronta para revisão', 'id': 'inicial',
         'data': fmt(minuta), 'responsavel': 'Adv. revisão'},
        {'marco': 'Protocolo da inicial revisada', 'id': 'protocolo',
         'data': fmt(minuta + timedelta(days=p['revisao_dias'])), 'responsavel': 'Protocolo'},
    ]


def entes(triagem):
    """Entes credores (Uniao/PGFN, Estado, Municipio, INSS) que a triagem encontrou nos debitos."""
    vistos = []
    for d in (triagem.get('debitos') or []):
        e = (d.get('ente') or '').strip()
        if e and e.lower() not in [x.lower() for x in vistos]:
            vistos.append(e)
    if not vistos:
        return ''
    return vistos[0] if len(vistos) == 1 else ', '.join(vistos[:-1]) + ' e ' + vistos[-1]


def proximo_numero_contrato():
    caminho = os.path.join(raiz_clientes(), '_controle_contratos.json')
    os.makedirs(os.path.dirname(caminho), exist_ok=True)
    try:
        with open(caminho, encoding='utf-8') as f:
            controle = json.load(f)
    except (FileNotFoundError, json.JSONDecodeError):
        controle = {}
    ano = str(date.today().year)
    controle[ano] = controle.get(ano, 0) + 1
    with open(caminho, 'w', encoding='utf-8') as f:
        json.dump(controle, f, indent=2)
    return f'{controle[ano]:03d}/{ano}'


def area_do_caso(triagem, cadastro):
    """PREVIDENCIARIO ou TRIBUTARIO. AMBOS e nao informado seguem TRIBUTARIO (carro-chefe)."""
    bruto = (cadastro.get('area') or triagem.get('area') or '').upper()
    return 'PREVIDENCIARIO' if bruto.startswith('PREV') else 'TRIBUTARIO'


def documentos_do_caso(area, cadastro):
    docs = list(DOCUMENTOS_POR_AREA[area])
    variante = (cadastro.get('contrato') or '').lower()
    if area == 'PREVIDENCIARIO' and 'sal' in variante:
        docs[0] = (docs[0][0], CONTRATO_EXITO_SALARIOS)
    if area == 'TRIBUTARIO' and 'resultado' in variante:
        docs[0] = (docs[0][0], CONTRATO_TRIB_COM_RESULTADO)
    if area == 'PREVIDENCIARIO' and (cadastro.get('juizado_federal') or '').lower() in ('sim', 's', 'true', '1'):
        docs.append(TERMO_RENUNCIA)
    return docs


def montar_campos(qualif, triagem, cadastro):
    h = triagem.get('honorarios') or {}
    ente = cadastro.get('ente') or entes(triagem)
    orgao = qualif.get('orgao_emissor', '')
    area = area_do_caso(triagem, cadastro)
    hp, ht = HONORARIOS['PREVIDENCIARIO'], HONORARIOS['TRIBUTARIO']
    deb = [d for d in (triagem.get('debitos') or []) if (d.get('identificacao') or '').strip()]
    campos = {
        'NOME': (qualif.get('nome') or '').upper(),
        'NACIONALIDADE': qualif.get('nacionalidade', ''),
        'ESTADO_CIVIL': qualif.get('estado_civil', ''),
        'PROFISSAO': qualif.get('profissao', ''),
        'RG': qualif.get('rg', ''),
        'ORGAO_EMISSOR': orgao,
        'CPF': qualif.get('cpf', ''),
        'ENDERECO_COMPLETO': endereco_completo(qualif),
        'TELEFONE': qualif.get('telefone', ''),
        'EMAIL': qualif.get('email', ''),
        'ENTE_CREDOR': ente,
        'OBJETO': cadastro.get('objeto') or (
            'Defesa dos direitos e interesses do(a) Outorgante em matéria tributária e/ou previdenciária'
            + (f', perante {ente}' if ente else '')
            + ', nas esferas administrativa e judicial, incluindo análise de débitos, defesas, recursos, '
            'embargos, exceção de pré-executividade, transação e parcelamento, e requerimentos junto ao INSS.'),
        'OUTORGADOS': '; '.join(f"{o['nome']}, {o['oab']}" for o in OUTORGADOS),
        'HONORARIOS_ENTRADA': cadastro.get('honorarios_entrada') or h.get('entrada', ''),
        'HONORARIOS_PAGAMENTO': cadastro.get('honorarios_pagamento') or h.get('parcelas', ''),
        'HONORARIOS_EXITO': cadastro.get('honorarios_exito') or h.get('exito', ''),
        'CIDADE_UF': f"{ESCRITORIO['cidade']}/{ESCRITORIO['uf']}",
        'DATA_EXTENSO': data_extenso(),
        'TITULAR_NOME': TITULAR['nome'].upper(),
        'TITULAR_OAB': TITULAR['oab'],
        'ESCRITORIO_CNPJ': ESCRITORIO['cnpj'],
        'ESCRITORIO_ENDERECO': ESCRITORIO['endereco'],
        'ESCRITORIO_EMAIL': ESCRITORIO['email'],
        'CIDADE': ESCRITORIO['cidade'],
        'DATA_NASCIMENTO': qualif.get('data_nascimento', ''),
        'TITULAR_EMAIL': TITULAR['email_contratos'],
        'TITULAR_ESTADO_CIVIL': TITULAR['estado_civil'],
        # previdenciario (Dra. Faya): politica de honorarios em config/escritorio.py
        'HONORARIOS_IMPLANTACAO': cadastro.get('honorarios_implantacao') or hp['implantacao'],
        'HONORARIOS_INICIAIS': cadastro.get('honorarios_iniciais') or hp['iniciais'],
        'OBJETO_ACAO': cadastro.get('objeto_acao') or cadastro.get('objeto') or '',
        # tributario (Dr. William): empresa contratante e dados do contrato
        'CNPJ': cadastro.get('cnpj') or qualif.get('cnpj', ''),
        'REPRESENTANTE_NOME': (cadastro.get('representante_nome') or qualif.get('nome') or '').upper(),
        'REPRESENTANTE_CPF': cadastro.get('representante_cpf') or qualif.get('cpf', ''),
        'REPRESENTANTE_QUALIFICACAO': cadastro.get('representante_qualificacao') or ', '.join(
            x for x in (qualif.get('nacionalidade'), qualif.get('estado_civil'), qualif.get('profissao'),
                        f"RG {qualif['rg']}" if qualif.get('rg') else '',
                        f"CPF {qualif['cpf']}" if qualif.get('cpf') else '') if x),
        'LISTA_CDAS': cadastro.get('lista_cdas') or ', '.join(d['identificacao'] for d in deb),
        'VALOR_DEBITO': cadastro.get('valor_debito') or '',
        'PRAZO_ADESAO': cadastro.get('prazo_adesao') or '',
        'HONORARIOS_CONDICOES': cadastro.get('honorarios_condicoes') or '',
        'CONTA_BANCARIA': os.getenv('CONTA_BANCARIA', ''),
        'CHAVE_PIX': os.getenv('CHAVE_PIX', ''),
    }
    if area == 'PREVIDENCIARIO':
        campos['OBJETO'] = cadastro.get('objeto') or ''
        campos['HONORARIOS_EXITO'] = campos['HONORARIOS_EXITO'] or hp['exito']
    return campos


def gerar(base, qualif, triagem, cadastro):
    campos = montar_campos(qualif, triagem, cadastro)
    campos['NUMERO_CONTRATO'] = proximo_numero_contrato()
    nome = nome_pasta(qualif.get('nome'))
    gerados = []
    area = area_do_caso(triagem, cadastro)
    for titulo_doc, modelo in documentos_do_caso(area, cadastro):
        caminho_modelo = os.path.join(MODELOS, modelo)
        if not os.path.exists(caminho_modelo):
            print(f'   AVISO: modelo {modelo} nao encontrado em DOCS_MODELOS/')
            continue
        saida = os.path.join(base, CONTRATACAO, f'{nome.title()} - {titulo_doc}.docx')
        faltando = preencher_modelo(caminho_modelo, campos, saida)
        gerados.append({'documento': titulo_doc, 'arquivo': saida, 'campos_faltando': sorted(faltando),
                        'pendencias': pendencias_no_arquivo(saida)})
    return gerados, campos


def pendencias_no_arquivo(caminho):
    """Tudo que ainda esta marcado [PREENCHER ...] ou [CONFERIR ...] no documento."""
    doc = Document(caminho)
    textos = [p.text for p in doc.paragraphs]
    for t in doc.tables:
        for linha in t.rows:
            textos.extend(c.text for c in linha.cells)
    achados = []
    for t in textos:
        achados.extend(PENDENCIA.findall(t))
    return sorted(set(achados))
