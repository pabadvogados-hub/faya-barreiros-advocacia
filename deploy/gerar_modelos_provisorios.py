"""
Gera os modelos PROVISORIOS de contrato, procuracao e declaracao em DOCS_MODELOS/.

Por que provisorios: o escritorio ainda nao mandou os modelos oficiais (briefing
pendente). Estes aqui tem a estrutura certa e os {{CAMPOS}} que o sistema preenche,
mas trazem [PREENCHER] nas clausulas que so o escritorio decide - e o
CONTRATACAO/envio.py TRAVA o envio enquanto existir [PREENCHER] ou [CONFERIR].

Assim da para rodar o fluxo inteiro desde hoje, sem risco de mandar ao cliente um
contrato com clausula que nao e do escritorio.

    python deploy/gerar_modelos_provisorios.py

Quando o escritorio mandar os modelos oficiais: salvar os .docx dele em DOCS_MODELOS/
com os mesmos nomes, trocando os dados variaveis pelos {{CAMPOS}} de
docs/CAMPOS_DOS_MODELOS.md. Nao precisa rodar este script de novo.
"""
import os
import sys

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.shared import Cm, Pt

RAIZ = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, RAIZ)
sys.path.insert(0, os.path.join(RAIZ, 'NUCLEO'))
import ambiente  # noqa: E402,F401
from config.escritorio import ESCRITORIO, VISUAL  # noqa: E402

DESTINO = os.path.join(RAIZ, 'DOCS_MODELOS')

AVISO = ('MODELO PROVISORIO - gerado pelo sistema enquanto o escritorio nao envia o modelo oficial. '
         'Os trechos marcados com [PREENCHER] travam o envio para assinatura de proposito.')


def _doc():
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    sec.top_margin, sec.bottom_margin = Cm(3), Cm(2.5)
    sec.left_margin, sec.right_margin = Cm(3), Cm(2.5)
    normal = doc.styles['Normal']
    normal.font.name = VISUAL['fonte']
    normal.font.size = Pt(VISUAL['tamanho'])
    pf = normal.paragraph_format
    pf.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    pf.space_after = Pt(6)
    return doc


def _p(doc, texto, negrito=False, centro=False, recuo=True):
    p = doc.add_paragraph()
    if centro:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    elif recuo:
        p.paragraph_format.first_line_indent = Cm(VISUAL['recuo_primeira_linha_cm'])
    r = p.add_run(texto)
    r.bold = negrito
    return p


def _titulo(doc, texto):
    _p(doc, texto.upper(), negrito=True, centro=True, recuo=False)


def _rodape_assinatura(doc):
    doc.add_paragraph()
    _p(doc, '{{CIDADE_UF}}, {{DATA_EXTENSO}}.', recuo=False)
    doc.add_paragraph()
    doc.add_paragraph()
    _p(doc, '_________________________________________', centro=True, recuo=False)
    _p(doc, '{{NOME}}', centro=True, recuo=False)
    _p(doc, 'CPF {{CPF}}', centro=True, recuo=False)
    doc.add_paragraph()
    _p(doc, '_________________________________________', centro=True, recuo=False)
    _p(doc, '{{TITULAR_NOME}}', centro=True, recuo=False)
    _p(doc, '{{TITULAR_OAB}}', centro=True, recuo=False)
    _p(doc, ESCRITORIO['razao_social'] + ' - CNPJ {{ESCRITORIO_CNPJ}}', centro=True, recuo=False)


def _cabecalho_aviso(doc):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(AVISO)
    r.bold = True
    r.font.size = Pt(8)


def contrato():
    doc = _doc()
    _cabecalho_aviso(doc)
    _titulo(doc, 'Contrato de prestação de serviços advocatícios')
    _p(doc, 'Contrato nº {{NUMERO_CONTRATO}}', recuo=False)
    doc.add_paragraph()
    _p(doc, 'CONTRATANTE: {{NOME}}, {{NACIONALIDADE}}, {{ESTADO_CIVIL}}, {{PROFISSAO}}, portador(a) da '
            'cédula de identidade nº {{RG}} {{ORGAO_EMISSOR}}, inscrito(a) no CPF sob o nº {{CPF}}, '
            'residente e domiciliado(a) em {{ENDERECO_COMPLETO}}, telefone {{TELEFONE}}, '
            'e-mail {{EMAIL}}.')
    _p(doc, f"CONTRATADO: {ESCRITORIO['razao_social']}, sociedade de advogados inscrita no CNPJ sob o nº "
            '{{ESCRITORIO_CNPJ}}, com sede em {{ESCRITORIO_ENDERECO}}, neste ato representada pelo '
            'advogado {{TITULAR_NOME}}, {{TITULAR_OAB}}.')
    doc.add_paragraph()
    _titulo(doc, 'Cláusula 1ª - do objeto')
    _p(doc, 'O CONTRATADO prestará ao CONTRATANTE serviços advocatícios em matéria tributária e/ou '
            'previdenciária, consistentes em [PREENCHER: análise da origem dos débitos / defesa administrativa / '
            'execução fiscal / transação / requerimento ou ação previdenciária], perante {{ENTE_CREDOR}}, '
            'bem como na prática dos atos necessários ao seu regular andamento nas esferas administrativa '
            'e judicial.')
    _titulo(doc, 'Cláusula 2ª - dos honorários')
    _p(doc, 'Honorários iniciais: {{HONORARIOS_ENTRADA}}.')
    _p(doc, 'Forma de pagamento: {{HONORARIOS_PAGAMENTO}}.')
    _p(doc, 'Honorários de êxito: {{HONORARIOS_EXITO}}, incidentes sobre o proveito econômico obtido, '
            'assim compreendido [PREENCHER: no tributário, se é o valor do débito reduzido ou extinto, a economia '
            'de juros e multa ou o valor restituído; no previdenciário, se incide sobre atrasados e/ou '
            'parcelas vincendas; e como ficam os honorários de sucumbência].')
    _p(doc, 'Os honorários de sucumbência pertencem ao CONTRATADO, nos termos do art. 23 da Lei 8.906/94.')
    _titulo(doc, 'Cláusula 3ª - das obrigações do contratante')
    _p(doc, 'O CONTRATANTE obriga-se a fornecer os documentos e as informações necessárias, a manter seus '
            'dados de contato atualizados e a comparecer aos atos para os quais for convocado.')
    _titulo(doc, 'Cláusula 4ª - das despesas processuais')
    _p(doc, '[PREENCHER: como o escritório trata custas, honorários periciais e despesas de deslocamento '
            'quando não há gratuidade de justiça, e de taxas e emolumentos administrativos.]')
    _titulo(doc, 'Cláusula 5ª - da rescisão')
    _p(doc, '[PREENCHER: regra de rescisão e de pagamento proporcional dos honorários em caso de revogação '
            'do mandato ou acordo feito diretamente pelo CONTRATANTE.]')
    _titulo(doc, 'Cláusula 6ª - do foro')
    _p(doc, 'Fica eleito o foro da comarca de {{CIDADE_UF}} para dirimir as questões oriundas deste contrato.')
    _rodape_assinatura(doc)
    return doc, 'CONTRATO_HONORARIOS_MODELO.docx'


def procuracao():
    doc = _doc()
    _cabecalho_aviso(doc)
    _titulo(doc, 'Procuração ad judicia et extra')
    doc.add_paragraph()
    _p(doc, 'OUTORGANTE: {{NOME}}, {{NACIONALIDADE}}, {{ESTADO_CIVIL}}, {{PROFISSAO}}, portador(a) da '
            'cédula de identidade nº {{RG}} {{ORGAO_EMISSOR}}, inscrito(a) no CPF sob o nº {{CPF}}, '
            'residente e domiciliado(a) em {{ENDERECO_COMPLETO}}.')
    _p(doc, 'OUTORGADOS: {{OUTORGADOS}}, todos integrantes de ' + ESCRITORIO['razao_social'] +
            ', CNPJ {{ESCRITORIO_CNPJ}}, com escritório em {{ESCRITORIO_ENDERECO}}.')
    doc.add_paragraph()
    _p(doc, 'PODERES: pelo presente instrumento, o OUTORGANTE nomeia e constitui seus bastantes procuradores '
            'os OUTORGADOS, a quem confere os poderes da cláusula ad judicia et extra, para o foro em geral, '
            'em qualquer juízo, instância ou tribunal, podendo propor contra quem de direito as ações '
            'competentes e defendê-lo nas contrárias, seguindo umas e outras até final decisão, usando dos '
            'recursos legais e acompanhando-os, conferindo-lhes, ainda, poderes especiais para confessar, '
            'desistir, transigir, firmar compromissos ou acordos, receber e dar quitação, firmar declaração '
            'de hipossuficiência econômica, substabelecer com ou sem reserva de poderes, agindo em conjunto '
            'ou separadamente.')
    _p(doc, 'OBJETO: {{OBJETO}}')
    _rodape_assinatura(doc)
    return doc, 'PROCURACAO_MODELO.docx'


def declaracao():
    doc = _doc()
    _cabecalho_aviso(doc)
    _titulo(doc, 'Declaração de hipossuficiência econômica')
    doc.add_paragraph()
    _p(doc, 'Eu, {{NOME}}, {{NACIONALIDADE}}, {{ESTADO_CIVIL}}, {{PROFISSAO}}, portador(a) da cédula de '
            'identidade nº {{RG}} {{ORGAO_EMISSOR}}, inscrito(a) no CPF sob o nº {{CPF}}, residente e '
            'domiciliado(a) em {{ENDERECO_COMPLETO}}, DECLARO, sob as penas da lei, não possuir condições '
            'de arcar com as custas processuais e os honorários advocatícios sem prejuízo do meu sustento '
            'e do de minha família, razão pela qual requeiro os benefícios da justiça gratuita, nos termos '
            'do art. 98 e seguintes do Código de Processo Civil.')
    _p(doc, 'Declaro, ainda, estar ciente de que a falsidade desta declaração sujeita o declarante às '
            'sanções civis, administrativas e criminais previstas na legislação.')
    doc.add_paragraph()
    _p(doc, '{{CIDADE_UF}}, {{DATA_EXTENSO}}.', recuo=False)
    doc.add_paragraph()
    doc.add_paragraph()
    _p(doc, '_________________________________________', centro=True, recuo=False)
    _p(doc, '{{NOME}}', centro=True, recuo=False)
    _p(doc, 'CPF {{CPF}}', centro=True, recuo=False)
    return doc, 'DECLARACAO_HIPOSSUFICIENCIA_MODELO.docx'


def main():
    os.makedirs(DESTINO, exist_ok=True)
    for montar in (contrato, procuracao, declaracao):
        doc, nome = montar()
        caminho = os.path.join(DESTINO, nome)
        if os.path.exists(caminho) and '--forcar' not in sys.argv:
            print(f'ja existe, nao sobrescrevi: {nome}  (use --forcar para refazer)')
            continue
        doc.save(caminho)
        print(f'gerado: {nome}')
    print('\nLembrete: estes modelos sao PROVISORIOS. O envio para assinatura fica travado '
          'enquanto houver [PREENCHER] neles.')


if __name__ == '__main__':
    main()
